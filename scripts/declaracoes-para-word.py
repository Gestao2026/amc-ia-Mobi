#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gerador dos modelos de declaração em papel timbrado.

Pega o conjunto de declarações escrito em markdown simples e gera o arquivo que
vai para o cliente: A4, timbrado repetido em toda página, uma declaração por
página, campos a preencher destacados em amarelo e bloco de assinatura.

  declaracoes-osc.md  ->  Declaracoes - OSC.docx
                          Declaracoes - OSC.pdf

Uso:
  python3 scripts/declaracoes-para-word.py <entrada.md> <saida.docx> [--sem-pdf]

Sintaxe do markdown de entrada:

  %%TIMBRADO ... %%FIM      bloco do cabeçalho, repetido em toda página
  %%ASSINATURA ... %%FIM    bloco usado onde aparecer o token [ASSINATURA]
  # Título                  capa, tudo antes da primeira declaração
  ## TÍTULO DA DECLARAÇÃO   começa uma declaração nova, em página nova
  :: texto                  parágrafo centralizado (assinatura de terceiro)
  > texto                   orientação interna, em caixa, para apagar antes de imprimir
  - item                    lista com marcador
  | tabela |                tabela real
  [ASSINATURA]              insere o bloco de assinatura
  {{campo}}                 campo a preencher, sai destacado em amarelo

Requer python-docx. O PDF depende do Word (Windows); sem ele, o .docx sai
normalmente e o PDF é ignorado com aviso.
"""

import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1F, 0x3B, 0x63)
CINZA = RGBColor(0x55, 0x55, 0x55)
CINZA_CLARO = "EDF1F7"
ORIENTACAO = "FFF6DA"

CAMPO = re.compile(r"(\{\{.+?\}\}|\*\*.+?\*\*)")


def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def add_bottom_border(paragraph, size=8, color="1F3B63"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def escrever(paragraph, texto, tamanho=11, cor=None, negrito=False, italico=False):
    """Escreve o texto resolvendo campos {{ }} em destaque e **negrito**."""
    for pedaco in CAMPO.split(texto):
        if not pedaco:
            continue
        bold = negrito
        destaque = False
        conteudo = pedaco
        if pedaco.startswith("{{") and pedaco.endswith("}}"):
            conteudo = pedaco[2:-2]
            destaque = True
        elif pedaco.startswith("**") and pedaco.endswith("**") and len(pedaco) > 4:
            conteudo = pedaco[2:-2]
            bold = True
        run = paragraph.add_run(conteudo)
        run.bold = bold
        run.italic = italico
        run.font.size = Pt(tamanho)
        if cor is not None:
            run.font.color.rgb = cor
        if destaque:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def ler_blocos(texto):
    """Separa os blocos %%NOME ... %%FIM do corpo do documento."""
    blocos = {}

    def captura(m):
        blocos[m.group(1)] = m.group(2).strip("\n")
        return ""

    corpo = re.sub(r"%%(\w+)\n(.*?)\n%%FIM", captura, texto, flags=re.S)
    return blocos, corpo


def montar_timbrado(section, linhas):
    """Monta o cabeçalho repetido em toda página."""
    header = section.header
    header.is_linked_to_previous = False
    for p in list(header.paragraphs):
        p._element.getparent().remove(p._element)

    for i, linha in enumerate(linhas):
        p = header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        if i == 0:
            escrever(p, linha, tamanho=10, cor=CINZA, italico=True)
        elif i == 1:
            escrever(p, linha, tamanho=13, cor=AZUL, negrito=True)
        else:
            escrever(p, linha, tamanho=8.5, cor=CINZA)
    if linhas:
        add_bottom_border(header.paragraphs[-1])
        header.paragraphs[-1].paragraph_format.space_after = Pt(8)


def montar_rodape(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for instr in ("begin", "instrText", "end"):
        el = OxmlElement("w:" + ("fldChar" if instr != "instrText" else "instrText"))
        if instr == "instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        else:
            el.set(qn("w:fldCharType"), instr)
        run._r.append(el)
    run.font.size = Pt(8)
    run.font.color.rgb = CINZA


def split_row(line):
    line = line.strip().strip("|")
    return [p.strip() for p in re.split(r"(?<!\\)\|", line)]


def is_separator(line):
    return bool(re.fullmatch(r"\|?[\s:\-|]+\|?", line.strip())) and "-" in line


def bloco_assinatura(doc, linhas):
    espaco = doc.add_paragraph()
    espaco.paragraph_format.space_before = Pt(24)
    espaco.paragraph_format.space_after = Pt(0)
    for linha in linhas:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        if linha.strip() == "":
            p.paragraph_format.space_after = Pt(14)
            continue
        escrever(p, linha, tamanho=11)


def build(md_path, out_path):
    with open(md_path, encoding="utf-8") as fh:
        texto = fh.read().replace("\r\n", "\n")

    blocos, corpo = ler_blocos(texto)
    timbrado = [l for l in blocos.get("TIMBRADO", "").split("\n")]
    assinatura = blocos.get("ASSINATURA", "").split("\n")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3.4)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.header_distance = Cm(1.2)
    sec.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.4

    montar_timbrado(sec, timbrado)
    montar_rodape(sec)

    linhas = corpo.split("\n")
    i = 0
    n = len(linhas)

    while i < n:
        linha = linhas[i].strip()

        if not linha:
            i += 1
            continue

        # nova declaração, em página nova
        m = re.match(r"^##\s+(.*)$", linha)
        if m:
            if doc.paragraphs or doc.tables:
                quebra = doc.add_paragraph()
                quebra.add_run().add_break(WD_BREAK.PAGE)
                quebra.paragraph_format.space_after = Pt(0)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(16)
            escrever(p, m.group(1).strip(), tamanho=12.5, cor=AZUL, negrito=True)
            i += 1
            continue

        # título da capa
        m = re.match(r"^#\s+(.*)$", linha)
        if m:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(14)
            escrever(p, m.group(1).strip(), tamanho=16, cor=AZUL, negrito=True)
            i += 1
            continue

        # subtítulo da capa
        m = re.match(r"^###\s+(.*)$", linha)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            escrever(p, m.group(1).strip(), tamanho=11.5, cor=AZUL, negrito=True)
            i += 1
            continue

        # bloco de assinatura
        if linha == "[ASSINATURA]":
            bloco_assinatura(doc, assinatura)
            i += 1
            continue

        # parágrafo centralizado
        if linha.startswith("::"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            escrever(p, linha[2:].strip(), tamanho=11)
            i += 1
            continue

        # orientação interna
        if linha.startswith(">"):
            bloco = []
            while i < n and linhas[i].strip().startswith(">"):
                bloco.append(linhas[i].strip().lstrip(">").strip())
                i += 1
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.1
            escrever(
                p,
                "Orientação, apagar antes de imprimir. " + " ".join(x for x in bloco if x),
                tamanho=9,
                cor=CINZA,
                italico=True,
            )
            set_cell_bg(cell, ORIENTACAO)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # tabela
        if linha.startswith("|") and i + 1 < n and is_separator(linhas[i + 1]):
            header = split_row(linha)
            i += 2
            rows = []
            while i < n and linhas[i].strip().startswith("|"):
                rows.append(split_row(linhas[i]))
                i += 1
            cols = max([len(header)] + [len(r) for r in rows])
            table = doc.add_table(rows=1, cols=cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            for c in range(cols):
                cell = hdr[c]
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                escrever(p, header[c] if c < len(header) else "", tamanho=9.5, cor=AZUL, negrito=True)
                set_cell_bg(cell, CINZA_CLARO)
            for r in rows:
                cells = table.add_row().cells
                for c in range(cols):
                    cell = cells[c]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.0
                    escrever(p, r[c] if c < len(r) else "", tamanho=9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # lista
        if re.match(r"^[-*+]\s+", linha):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.2
            escrever(p, re.sub(r"^[-*+]\s+", "", linha))
            i += 1
            continue

        if re.match(r"^\d+[.)]\s+", linha):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.2
            escrever(p, re.sub(r"^\d+[.)]\s+", "", linha))
            i += 1
            continue

        # parágrafo comum
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        escrever(p, linha)
        i += 1

    doc.save(out_path)
    print("OK ->", out_path)


def gerar_pdf(docx_path):
    """Converte o .docx em PDF usando o Word instalado (somente Windows)."""
    import os
    import subprocess

    docx_abs = os.path.abspath(docx_path)
    pdf_abs = os.path.splitext(docx_abs)[0] + ".pdf"
    ps = (
        "$w = New-Object -ComObject Word.Application; "
        "$w.Visible=$false; $w.DisplayAlerts=0; "
        "$d = $w.Documents.Open('%s',$false,$true); "
        "$d.SaveAs([ref]'%s',[ref]17); "
        "$d.Close([ref]0); $w.Quit()" % (docx_abs, pdf_abs)
    )
    try:
        subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
            check=True,
            capture_output=True,
            timeout=240,
        )
    except Exception as erro:
        print("PDF nao gerado (Word indisponivel):", erro)
        return None
    print("OK ->", pdf_abs)
    return pdf_abs


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) < 2:
        print(__doc__)
        sys.exit(1)
    build(args[0], args[1])
    if "--sem-pdf" not in sys.argv:
        gerar_pdf(args[1])
