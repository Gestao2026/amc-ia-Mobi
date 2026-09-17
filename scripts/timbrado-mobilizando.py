#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Conversor de markdown para documento no papel timbrado da Mobilizando.

Serve para todo material que sai da assessoria para o cliente: relação de
documentos, orientações, roteiros, relatórios. Logomarca no cabeçalho, régua
dourada, rodapé com endereço e contato, numeração de página.

  documentos-osc.md  ->  Documentos - OSC.docx
                         Documentos - OSC.pdf

Uso:
  python3 scripts/timbrado-mobilizando.py <entrada.md> <saida.docx> [--sem-pdf]

Suporta: # ## ### títulos, parágrafos, listas, listas de conferência (- [ ]),
tabelas, blocos de destaque (>), régua horizontal (---), **negrito** e os
selos `[base]`, `[edital]` e `[pós]`.

Requer python-docx. O PDF depende do Word (Windows); sem ele, o .docx sai
normalmente e o PDF é ignorado com aviso.
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x0F, 0x2B, 0x6C)
DOURADO = RGBColor(0x8A, 0x6D, 0x2F)
CINZA = RGBColor(0x55, 0x55, 0x55)
CINZA_CLARO = "E8EDF6"
DESTAQUE = "FBF6E9"
LINHA_AZUL = "0F2B6C"
LINHA_DOURADA = "D1B484"

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO = os.path.join(RAIZ, "modelos", "marca", "logo-mobilizando.png")

ASSINATURA_RODAPE = "Mobilizando. Assessoria em Captação de Recursos"
ENDERECO_RODAPE = "Rua Arcos, 643, Vera Cruz, Belo Horizonte, MG, CEP 30.285-100"
CONTATO_RODAPE = (
    "+55 (31) 97150-5768   |   rosepaula@mobilizando.org   |   @mobilizando_impactosocial"
)

SELOS = ("[base]", "[edital]", "[pós]", "[pos]")

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[ x]\]|\[.+?\]\(.+?\))")


def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def add_border(paragraph, lado="bottom", size=6, color=LINHA_AZUL):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = pPr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        pPr.append(pbdr)
    el = OxmlElement("w:" + lado)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "3")
    el.set(qn("w:color"), color)
    pbdr.append(el)


def cor_do_selo(texto):
    t = texto.lower()
    if "base" in t:
        return AZUL
    if "edital" in t:
        return DOURADO
    return CINZA


def add_runs(paragraph, texto, base_size=None, base_color=None, base_bold=False):
    """Escreve o texto com negrito, selos, caixas de conferência e links."""
    texto = texto.replace("\\|", "|")
    for pedaco in INLINE.split(texto):
        if not pedaco:
            continue
        bold = base_bold
        cor = base_color
        tamanho = base_size or 10.5
        conteudo = pedaco
        fonte = None

        if pedaco.startswith("**") and pedaco.endswith("**") and len(pedaco) > 4:
            conteudo, bold = pedaco[2:-2], True
        elif pedaco.startswith("`") and pedaco.endswith("`") and len(pedaco) > 2:
            interno = pedaco[1:-1]
            if interno.lower() in SELOS:
                conteudo = interno.strip("[]").upper()
                bold = True
                tamanho = 8
                cor = cor_do_selo(interno)
            else:
                conteudo = interno
                fonte = "Consolas"
                tamanho = (base_size or 10.5) - 0.5
        elif pedaco in ("[ ]", "[x]"):
            conteudo = "\u2612  " if pedaco == "[x]" else "\u2610  "
            fonte = "Segoe UI Symbol"
            tamanho = (base_size or 10.5) + 1
            cor = AZUL
        else:
            m = re.fullmatch(r"\[(.+?)\]\((.+?)\)", pedaco)
            if m:
                rotulo, url = m.group(1), m.group(2)
                conteudo = rotulo if rotulo == url else "%s (%s)" % (rotulo, url)

        run = paragraph.add_run(conteudo)
        run.bold = bold
        run.font.size = Pt(tamanho)
        if fonte:
            run.font.name = fonte
        if cor is not None:
            run.font.color.rgb = cor


def montar_timbrado(section):
    header = section.header
    header.is_linked_to_previous = False
    for p in list(header.paragraphs):
        p._element.getparent().remove(p._element)

    p = header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    if os.path.exists(LOGO):
        p.add_run().add_picture(LOGO, width=Cm(5.6))
    else:
        add_runs(p, "Mobilizando", base_size=16, base_color=AZUL, base_bold=True)

    regua = header.add_paragraph()
    regua.paragraph_format.space_before = Pt(0)
    regua.paragraph_format.space_after = Pt(6)
    add_border(regua, "bottom", size=12, color=LINHA_DOURADA)


def montar_rodape(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)

    topo = footer.add_paragraph()
    topo.paragraph_format.space_before = Pt(0)
    topo.paragraph_format.space_after = Pt(3)
    add_border(topo, "bottom", size=6, color=LINHA_DOURADA)

    p1 = footer.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(0)
    add_runs(p1, ASSINATURA_RODAPE, base_size=8, base_color=AZUL, base_bold=True)

    for texto in (ENDERECO_RODAPE, CONTATO_RODAPE):
        p2 = footer.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.0
        add_runs(p2, texto, base_size=7.5, base_color=CINZA)

    p3 = footer.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(2)
    run = p3.add_run()
    for instr in ("begin", "instrText", "end"):
        el = OxmlElement("w:" + ("fldChar" if instr != "instrText" else "instrText"))
        if instr == "instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        else:
            el.set(qn("w:fldCharType"), instr)
        run._r.append(el)
    run.font.size = Pt(8)
    run.font.color.rgb = CINZA


def split_row(line):
    line = line.strip().strip("|")
    return [p.strip() for p in re.split(r"(?<!\\)\|", line)]


def is_separator(line):
    return bool(re.fullmatch(r"\|?[\s:\-|]+\|?", line.strip())) and "-" in line


def build(md_path, out_path):
    with open(md_path, encoding="utf-8") as fh:
        linhas = fh.read().replace("\r\n", "\n").split("\n")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3.3)
    sec.bottom_margin = Cm(2.6)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(1.1)
    sec.footer_distance = Cm(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    montar_timbrado(sec)
    montar_rodape(sec)

    i = 0
    n = len(linhas)

    while i < n:
        raw = linhas[i]
        linha = raw.strip()

        if not linha:
            i += 1
            continue

        # régua horizontal
        if re.fullmatch(r"-{3,}|\*{3,}", linha):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            add_border(p, "bottom", size=4, color="C9D2E0")
            i += 1
            continue

        # tabela
        if linha.startswith("|") and i + 1 < n and is_separator(linhas[i + 1]):
            header = split_row(linha)
            i += 2
            rows = []
            while i < n and linhas[i].strip().startswith("|"):
                rows.append(split_row(linhas[i]))
                i += 1
            cols = max([len(header)] + [len(r) for r in rows])
            table = doc.add_table(rows=1, cols=cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr = table.rows[0].cells
            for c in range(cols):
                cell = hdr[c]
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                add_runs(p, header[c] if c < len(header) else "", base_size=9.5,
                         base_color=AZUL, base_bold=True)
                set_cell_bg(cell, CINZA_CLARO)
            for r in rows:
                cells = table.add_row().cells
                for c in range(cols):
                    cell = cells[c]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.line_spacing = 1.0
                    add_runs(p, r[c] if c < len(r) else "", base_size=9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # bloco de destaque
        if linha.startswith(">"):
            bloco = []
            while i < n and linhas[i].strip().startswith(">"):
                bloco.append(linhas[i].strip().lstrip(">").strip())
                i += 1
            texto = " ".join(x for x in bloco if x)
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.rows[0].cells[0]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, texto, base_size=10)
            set_cell_bg(cell, DESTAQUE)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # títulos
        m = re.match(r"^(#{1,4})\s+(.*)$", linha)
        if m:
            nivel = len(m.group(1))
            texto = m.group(2).strip()
            p = doc.add_paragraph()
            pf = p.paragraph_format
            if nivel == 1:
                pf.space_before = Pt(4)
                pf.space_after = Pt(10)
                add_runs(p, texto, base_size=17, base_color=AZUL, base_bold=True)
                add_border(p, "bottom", size=8, color=LINHA_AZUL)
            elif nivel == 2:
                pf.space_before = Pt(16)
                pf.space_after = Pt(6)
                add_runs(p, texto, base_size=13.5, base_color=AZUL, base_bold=True)
                add_border(p, "bottom", size=4, color=LINHA_DOURADA)
            elif nivel == 3:
                pf.space_before = Pt(10)
                pf.space_after = Pt(4)
                add_runs(p, texto, base_size=11.5, base_color=AZUL, base_bold=True)
            else:
                pf.space_before = Pt(8)
                pf.space_after = Pt(3)
                add_runs(p, texto, base_size=11, base_color=CINZA, base_bold=True)
            i += 1
            continue

        # lista de conferência
        m = re.match(r"^[-*+]\s+(\[[ x]\]\s*)(.*)$", linha)
        if m:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.8)
            pf.first_line_indent = Cm(-0.8)
            pf.space_after = Pt(4)
            pf.line_spacing = 1.1
            add_runs(p, m.group(1).strip() + m.group(2))
            i += 1
            continue

        # lista com marcador
        if re.match(r"^[-*+]\s+", linha):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, re.sub(r"^[-*+]\s+", "", linha))
            i += 1
            continue

        # lista numerada
        if re.match(r"^\d+[.)]\s+", linha):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, re.sub(r"^\d+[.)]\s+", "", linha))
            i += 1
            continue

        # parágrafo comum
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, linha)
        i += 1

    doc.save(out_path)
    print("OK ->", out_path)


def gerar_pdf(docx_path):
    """Converte o .docx em PDF usando o Word instalado (somente Windows)."""
    import subprocess

    docx_abs = os.path.abspath(docx_path)
    pdf_abs = os.path.splitext(docx_abs)[0] + ".pdf"
    ps = (
        "$w = New-Object -ComObject Word.Application; "
        "$w.Visible=$false; $w.DisplayAlerts=0; "
        "$d = $w.Documents.Open('%s',$false,$true); "
        "$d.SaveAs([ref]'%s',[ref]17); "
        "$d.Close([ref]0); $w.Quit()" % (docx_abs, pdf_abs)
    )
    try:
        subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
            check=True,
            capture_output=True,
            timeout=240,
        )
    except Exception as erro:
        print("PDF nao gerado (Word indisponivel):", erro)
        return None
    print("OK ->", pdf_abs)
    return pdf_abs


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) < 2:
        print(__doc__)
        sys.exit(1)
    build(args[0], args[1])
    if "--sem-pdf" not in sys.argv:
        gerar_pdf(args[1])
